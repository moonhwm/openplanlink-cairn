#!/usr/bin/env python
# -*- coding: utf-8 -*-
# BLAST: NONE —— 本件只读本机 .docx，不触网。
r"""probe_meow_constraints.py —— 从 74 号令原件里取 MeoW 的【原文约束】。 v1.0.0
为何：轮133 我说协调稿缺「限流（3 秒/15/60/1000）」——**而我自己的 outbox 里搜不到这组数字**
⇒ 故先核【我到底有没有记下】，再谈别人缺不缺。
"""
from __future__ import annotations
import pathlib
import re
import sys

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:
        pass

DESK = pathlib.Path(r"C:\Users\欧阳宏俊\OneDrive\桌面")
cands = list(DESK.glob("*74号令*.docx")) + list(DESK.glob("*74*.docx"))
if not cands:
    print("★未找到 74号令 .docx")
    sys.exit(1)
p = cands[0]
print("═══ 读 74 号令原件：%s（%d B）═══" % (p.name, p.stat().st_size))

try:
    from docx import Document
except Exception as e:
    print("★缺 python-docx：%s" % e)
    sys.exit(2)

doc = Document(str(p))
paras = [x.text for x in doc.paragraphs]
tbl_txt = []
for t in doc.tables:
    for r in t.rows:
        tbl_txt.append(" | ".join(c.text.replace("\n", " ").strip() for c in r.cells))
allt = paras + tbl_txt
print("  段落 %d ｜ 表格行 %d" % (len(paras), len(tbl_txt)))

PAT = re.compile(r"MeoW|限流|3\s*秒|15\s*(?:条|次)?\s*/?\s*(?:分|分钟)|60\s*(?:条|次)?\s*/?\s*(?:时|小时)|1000|昵称|api\.chuckfang|凭据", re.I)
hits = [(i, s.strip()) for i, s in enumerate(allt) if PAT.search(s)]
print("  命中 %d 处：" % len(hits))
for i, s in hits[:40]:
    print("     [%d] %s" % (i, s[:150]))
print()
# 明确回答那个问题
num = re.compile(r"3\s*秒|15\s*(?:条|次)?\s*/?\s*(?:分|分钟)|60\s*(?:条|次)?\s*/?\s*(?:时|小时)|1000\s*(?:条|次)?\s*/?\s*(?:日|天)")
found = [s for _, s in hits if num.search(s)]
print("  ── 判词 ──")
if found:
    print("     ★ 74 号令原件里【确有限流数字】：")
    for s in found[:8]:
        print("        " + s[:140])
else:
    print("     ★ 74 号令原件里【未见】那组限流数字（3 秒／15 分／60 时／1000 日）")
    print("     ⇒ 若如此，则轮133 我说的『协调稿缺限流』这句话【依据不足】——")
    print("       因为**我自己也没有把该组数字记在台账或出件里** ⇒ 须更正。")
