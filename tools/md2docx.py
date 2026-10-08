#!/usr/bin/env python
# -*- coding: utf-8 -*-
r"""
md2docx.py —— 极简 Markdown → .docx 转换器（面向本席预印稿体例）。   v1.0.0

为什么自己写
────────────
《2026-9-25-OpenPlanLink 润色-1 (3)》第 6 条要求成果**以完整程序形式**交付；
第 8 条又要求研究进展**以 arXiv 论文格式**撰文。
本环境**无 LaTeX／pandoc／PDF 链**，但有 `python-docx`。
⇒ 故把「Markdown 正本 → docx 可开件」这一步**写成程序**：既产出体例件，又使这一步本身可复现。

支持子集（够用即止，多一项就多一处会错）
────────────────────────────────────────
  `# ` / `## ` / `### `   标题（1/2/3 级）
  `| a | b |` 连续行      表格（首行为表头；`|---|` 分隔行自动跳过）
  `- `                   无序列表
  `1. `                  有序列表
  ``` 围栏              等宽段落
  `**粗体**`             行内粗体
  `---` 独占一行          跳过（作分节留白）
  其余非空行              普通段落

用法：md2docx.py <输入.md> <输出.docx>
零第三方依赖之外仅需 python-docx。
"""
from __future__ import annotations
import pathlib
import re
import sys

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:
        pass

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
except Exception as e:                                  # noqa: BLE001
    print("[ERR] 需要 python-docx：%s" % e)
    sys.exit(2)

BOLD = re.compile(r"\*\*(.+?)\*\*")


def add_rich(p, text: str):
    """把 **粗体** 拆成 run，其余按普通文本。"""
    pos = 0
    for m in BOLD.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        r = p.add_run(m.group(1))
        r.bold = True
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])


def set_cjk_font(doc):
    """正文字体设为中西兼容，避免中文方框。"""
    try:
        st = doc.styles["Normal"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    except Exception:                                   # noqa: BLE001
        pass


def convert(src: pathlib.Path, dst: pathlib.Path) -> dict:
    lines = src.read_text(encoding="utf-8").splitlines()
    doc = Document()
    set_cjk_font(doc)
    stats = {"heading": 0, "para": 0, "bullet": 0, "number": 0, "table": 0, "code": 0}
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        if s.startswith("```"):                          # 围栏代码块
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"
            stats["code"] += 1
            continue

        if s.startswith("|") and s.endswith("|"):        # 表格
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                t.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, val in enumerate(row):
                        if ci < len(t.rows[ri].cells):
                            cell = t.rows[ri].cells[ci]
                            cell.text = ""
                            add_rich(cell.paragraphs[0], val)
                            if ri == 0:
                                for rr in cell.paragraphs[0].runs:
                                    rr.bold = True
                stats["table"] += 1
            continue

        if s.startswith("### "):
            doc.add_heading(s[4:], level=3); stats["heading"] += 1
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2); stats["heading"] += 1
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1); stats["heading"] += 1
        elif s == "---" or not s:
            pass
        elif s.startswith("- "):
            add_rich(doc.add_paragraph(style="List Bullet"), s[2:]); stats["bullet"] += 1
        elif re.match(r"^\d+\.\s", s):
            add_rich(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s", "", s)); stats["number"] += 1
        else:
            add_rich(doc.add_paragraph(), s); stats["para"] += 1
        i += 1

    doc.save(str(dst))
    return stats


def main():
    if len(sys.argv) < 3:
        print("用法：md2docx.py <输入.md> <输出.docx>")
        return 2
    src, dst = pathlib.Path(sys.argv[1]), pathlib.Path(sys.argv[2])
    if not src.is_file():
        print("[ERR] 输入不存在：%s" % src); return 2
    st = convert(src, dst)
    print("[OK] 已转换 → %s（%d B）" % (dst, dst.stat().st_size))
    print("     标题 %d ｜ 段 %d ｜ 无序 %d ｜ 有序 %d ｜ 表 %d ｜ 代码块 %d"
          % (st["heading"], st["para"], st["bullet"], st["number"], st["table"], st["code"]))
    print("     ★本步只做体例转换，不改一字内容；Markdown 正本仍为唯一权威版本。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
